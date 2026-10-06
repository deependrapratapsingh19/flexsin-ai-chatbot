from io import BytesIO

import pandas as pd
from docx import Document
from pypdf import PdfReader


# ============================================================
# CONFIGURATION
# ============================================================

SUPPORTED_DOCUMENT_TYPES = {
    "pdf",
    "docx",
    "txt",
    "csv",
}


# Large documents ko unlimited prompt me bhejne se bachane ke liye
# Version 1 me extracted context limit rakhenge.
MAX_EXTRACTED_CHARACTERS = 60000

# CSV me maximum rows jo direct LLM context me jayengi
MAX_CSV_ROWS = 200


# ============================================================
# FILE EXTENSION
# ============================================================

def get_file_extension(uploaded_file):

    if uploaded_file is None:
        return ""

    filename = getattr(
        uploaded_file,
        "name",
        ""
    )

    if "." not in filename:
        return ""

    return filename.rsplit(
        ".",
        1
    )[-1].lower().strip()


# ============================================================
# PDF PROCESSOR
# ============================================================

def extract_pdf_text(uploaded_file):

    try:

        file_bytes = uploaded_file.getvalue()

        pdf_stream = BytesIO(
            file_bytes
        )

        reader = PdfReader(
            pdf_stream
        )

        extracted_pages = []


        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            try:

                page_text = page.extract_text()

            except Exception:

                page_text = ""


            if page_text:

                page_text = page_text.strip()


                if page_text:

                    extracted_pages.append(
                        f"""
--- Page {page_number} ---

{page_text}
""".strip()
                    )


        if not extracted_pages:

            return (
                "No readable text could be extracted "
                "from this PDF. The PDF may be scanned "
                "or image-based."
            )


        complete_text = "\n\n".join(
            extracted_pages
        )


        return limit_text(
            complete_text
        )


    except Exception as error:

        raise ValueError(
            f"PDF read nahi ho paya: {error}"
        ) from error


# ============================================================
# DOCX PROCESSOR
# ============================================================

def extract_docx_text(uploaded_file):

    try:

        file_bytes = uploaded_file.getvalue()

        docx_stream = BytesIO(
            file_bytes
        )

        document = Document(
            docx_stream
        )


        extracted_content = []


        # ====================================================
        # PARAGRAPHS
        # ====================================================

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                extracted_content.append(
                    text
                )


        # ====================================================
        # TABLES
        # ====================================================

        for table_number, table in enumerate(
            document.tables,
            start=1
        ):

            table_rows = []


            for row in table.rows:

                row_values = []


                for cell in row.cells:

                    cell_text = (
                        cell.text
                        .replace("\n", " ")
                        .strip()
                    )

                    row_values.append(
                        cell_text
                    )


                table_rows.append(
                    " | ".join(
                        row_values
                    )
                )


            if table_rows:

                table_text = "\n".join(
                    table_rows
                )


                extracted_content.append(
                    f"""
--- Table {table_number} ---

{table_text}
""".strip()
                )


        if not extracted_content:

            return (
                "No readable text was found "
                "inside this DOCX file."
            )


        complete_text = "\n\n".join(
            extracted_content
        )


        return limit_text(
            complete_text
        )


    except Exception as error:

        raise ValueError(
            f"DOCX file read nahi ho payi: {error}"
        ) from error


# ============================================================
# TXT PROCESSOR
# ============================================================

def extract_txt_text(uploaded_file):

    try:

        file_bytes = uploaded_file.getvalue()


        # Common encodings ko try karo
        encodings = [
            "utf-8",
            "utf-8-sig",
            "cp1252",
            "latin-1",
        ]


        decoded_text = None


        for encoding in encodings:

            try:

                decoded_text = file_bytes.decode(
                    encoding
                )

                break

            except UnicodeDecodeError:

                continue


        if decoded_text is None:

            raise ValueError(
                "Text encoding identify nahi ho payi."
            )


        decoded_text = decoded_text.strip()


        if not decoded_text:

            return "TXT file empty hai."


        return limit_text(
            decoded_text
        )


    except Exception as error:

        raise ValueError(
            f"TXT file read nahi ho payi: {error}"
        ) from error


# ============================================================
# CSV PROCESSOR
# ============================================================

def extract_csv_text(uploaded_file):

    try:

        file_bytes = uploaded_file.getvalue()


        # ====================================================
        # FIRST TRY NORMAL UTF-8
        # ====================================================

        try:

            dataframe = pd.read_csv(
                BytesIO(file_bytes)
            )

        except UnicodeDecodeError:

            dataframe = pd.read_csv(
                BytesIO(file_bytes),
                encoding="latin-1"
            )


        if dataframe.empty:

            return "CSV file me koi data nahi hai."


        total_rows = len(
            dataframe
        )

        total_columns = len(
            dataframe.columns
        )


        # ====================================================
        # COLUMN INFORMATION
        # ====================================================

        column_names = [
            str(column)
            for column
            in dataframe.columns
        ]


        # ====================================================
        # LIMIT ROWS SENT TO LLM
        # ====================================================

        preview_dataframe = dataframe.head(
            MAX_CSV_ROWS
        )


        csv_text = preview_dataframe.to_csv(
            index=False
        )


        metadata = f"""
CSV INFORMATION

Total rows: {total_rows}
Total columns: {total_columns}

Columns:
{", ".join(column_names)}

DATA
{csv_text}
""".strip()


        if total_rows > MAX_CSV_ROWS:

            metadata += (
                f"\n\nNote: CSV contains {total_rows} rows. "
                f"Only the first {MAX_CSV_ROWS} rows are "
                "included in the current direct-analysis context."
            )


        return limit_text(
            metadata
        )


    except Exception as error:

        raise ValueError(
            f"CSV file read nahi ho payi: {error}"
        ) from error


# ============================================================
# LIMIT EXTRACTED TEXT
# ============================================================

def limit_text(text):

    if not text:

        return ""


    if len(text) <= MAX_EXTRACTED_CHARACTERS:

        return text


    truncated_text = text[
        :MAX_EXTRACTED_CHARACTERS
    ]


    return (
        truncated_text
        + "\n\n"
        + "[Document content truncated because the file "
        + "is too large for direct analysis. "
        + "A RAG pipeline should be used for full large-document "
        + "question answering.]"
    )


# ============================================================
# MAIN DOCUMENT PROCESSOR
# ============================================================

def process_document(uploaded_file):

    if uploaded_file is None:

        raise ValueError(
            "No file was provided."
        )


    extension = get_file_extension(
        uploaded_file
    )


    if extension not in SUPPORTED_DOCUMENT_TYPES:

        raise ValueError(
            f"Unsupported document type: .{extension}"
        )


    # ========================================================
    # PDF
    # ========================================================

    if extension == "pdf":

        text = extract_pdf_text(
            uploaded_file
        )


    # ========================================================
    # DOCX
    # ========================================================

    elif extension == "docx":

        text = extract_docx_text(
            uploaded_file
        )


    # ========================================================
    # TXT
    # ========================================================

    elif extension == "txt":

        text = extract_txt_text(
            uploaded_file
        )


    # ========================================================
    # CSV
    # ========================================================

    elif extension == "csv":

        text = extract_csv_text(
            uploaded_file
        )


    else:

        raise ValueError(
            "Unsupported document type."
        )


    return {
        "name": uploaded_file.name,
        "extension": extension,
        "text": text,
    }


# ============================================================
# PROCESS MULTIPLE DOCUMENTS
# ============================================================

def process_documents(uploaded_files):

    processed_documents = []


    if not uploaded_files:

        return processed_documents


    for uploaded_file in uploaded_files:

        extension = get_file_extension(
            uploaded_file
        )


        if extension not in SUPPORTED_DOCUMENT_TYPES:

            continue


        processed_document = process_document(
            uploaded_file
        )


        processed_documents.append(
            processed_document
        )


    return processed_documents


# ============================================================
# COMBINE DOCUMENT CONTEXT
# ============================================================

def build_document_context(processed_documents):

    if not processed_documents:

        return ""


    context_parts = []


    for document in processed_documents:

        context_parts.append(
            f"""
==============================
FILE: {document["name"]}
TYPE: {document["extension"].upper()}
==============================

{document["text"]}
""".strip()
        )


    combined_context = "\n\n".join(
        context_parts
    )


    return limit_text(
        combined_context
    )
